from pathlib import Path
from docx import Document


def cell_text(cell):
    return " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip()).replace("|", "\\|")


def extract(source: Path) -> None:
    document = Document(source)
    lines = [f"# Extract of {source.name}", "", "> Mechanically extracted from the authoritative DOCX; paragraphs and tables are preserved in document order where possible.", ""]
    table_by_element = {table._element: table for table in document.tables}
    paragraph_by_element = {paragraph._element: paragraph for paragraph in document.paragraphs}

    for element in document.element.body.iterchildren():
        if element in paragraph_by_element:
            paragraph = paragraph_by_element[element]
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            if style.startswith("Heading"):
                try:
                    level = min(6, int(style.split()[-1]))
                except ValueError:
                    level = 2
                lines.extend([f"{'#' * level} {text}", ""])
            else:
                lines.extend([text, ""])
        elif element in table_by_element:
            table = table_by_element[element]
            rows = [[cell_text(cell) for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("| " + " | ".join(["---"] * width) + " |")
            for row in rows[1:]:
                lines.append("| " + " | ".join(row) + " |")
            lines.append("")

    output = source.with_suffix(".extracted.md")
    output.write_text("\n".join(lines), encoding="utf-8")
    print(f"{source.name}: {len(document.paragraphs)} paragraphs, {len(document.tables)} tables -> {output.name}")


if __name__ == "__main__":
    for path in sorted(Path("docs/source").glob("*.docx")):
        extract(path)
